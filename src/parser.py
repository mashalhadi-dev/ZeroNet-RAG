"""Fault-tolerant corpus walking and text extraction."""
from __future__ import annotations

import csv
import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

log = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    path: str
    text: str


def _read_text(path: Path) -> str:
    # Try common encodings without allowing a bad byte to abort ingestion.
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(encoding)
        except (UnicodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def _extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rst", ".log", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".py", ".toml", ".ini", ".cfg"}:
        return _read_text(path)
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ValueError("encrypted PDF")
            except Exception as exc:
                raise ValueError("encrypted PDF") from exc
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(parts)
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            return "\n".join(" | ".join(row) for row in csv.reader(handle))
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        lines = []
        try:
            for sheet in workbook.worksheets:
                lines.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    vals = [str(value) for value in row if value is not None]
                    if vals:
                        lines.append(" | ".join(vals))
        finally:
            workbook.close()
        return "\n".join(lines)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        try:
            from PIL import Image
            import pytesseract
            with Image.open(path) as image:
                return pytesseract.image_to_string(image)
        except ImportError as exc:
            raise RuntimeError("image OCR requires Pillow and pytesseract") from exc
    raise ValueError(f"unsupported file type: {suffix or '(no extension)'}")


def walk_corpus(corpus_dir: str | Path) -> Iterator[ParsedDocument]:
    """Yield readable documents; each file failure is logged and skipped."""
    root = Path(corpus_dir).resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"corpus directory does not exist: {root}")
    for current, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = sorted(d for d in dirs if not (Path(current) / d).is_symlink())
        for filename in sorted(files):
            path = Path(current) / filename
            try:
                # Check mode bits explicitly because containers may run as root,
                # which can otherwise read files marked chmod 000.
                if path.is_symlink() or path.stat().st_mode & 0o444 == 0:
                    raise PermissionError("file has no read permission bits")
                text = _extract(path)
                if text and text.strip():
                    yield ParsedDocument(path.relative_to(root).as_posix(), text)
                else:
                    log.warning("Skipping empty document: %s", path)
            except Exception as exc:  # One bad file must not stop corpus ingestion.
                log.warning("Skipping %s (%s: %s)", path, type(exc).__name__, exc)
